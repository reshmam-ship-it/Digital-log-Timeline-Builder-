
import csv
import json
import sys
import queue
import threading
import tkinter as tk
import xml.etree.ElementTree as ET

from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from tkinter import ttk, filedialog, messagebox


# Application title and interface colors
APP_TITLE = "FORENSIGHT | Digital Forensics Log Timeline Builder"

BG = "#0b1020"
PANEL = "#111a2e"
PANEL2 = "#17233b"
BORDER = "#263653"
TEXT = "#e8eefc"
MUTED = "#8fa2c7"
ACCENT = "#54d6ff"
GREEN = "#57e39b"
RED = "#ff6b81"
PURPLE = "#b7a0ff"


# Store the details of each log event
@dataclass
class LogEvent:
    timestamp: str
    event_type: str
    source: str
    description: str
    host: str = ""
    user: str = ""
    source_ip: str = ""
    severity: str = "Info"
    event_id: str = ""
    event_id_internal: int = 0


# Each node stores an event and links to the next node
class Node:
    def __init__(self, event):
        self.event = event
        self.next = None


# Manage log events using a singly linked list
class EventLinkedList:
    def __init__(self):
        self.head = None
        self.tail = None
        self.size = 0

    # Add an event to the end of the linked list
    def append(self, event):
        node = Node(event)

        if self.head is None:
            self.head = self.tail = node
        else:
            self.tail.next = node
            self.tail = node

        self.size += 1

    # Visit each node in the linked list
    def traverse(self):
        current = self.head

        while current:
            yield current.event
            current = current.next

    # Search for events using a keyword
    def find(self, keyword=""):
        keyword = keyword.casefold().strip()
        events = list(self.traverse())

        if not keyword:
            return events

        return [
            event for event in events
            if keyword in " ".join(
                str(value) for value in asdict(event).values()
            ).casefold()
        ]

    # Delete an event using its internal ID
    def delete_by_id(self, internal_id):
        current = self.head
        previous = None

        while current:
            if current.event.event_id_internal == internal_id:
                if previous is None:
                    self.head = current.next
                else:
                    previous.next = current.next

                if current is self.tail:
                    self.tail = previous

                self.size -= 1
                return True

            previous = current
            current = current.next

        return False

    # Remove all events from the linked list
    def clear(self):
        self.head = self.tail = None
        self.size = 0


# Convert a timestamp into a standard datetime format
def parse_datetime(value):
    if not value:
        return None

    value = str(value).strip()

    try:
        return datetime.fromisoformat(
            value.replace("Z", "+00:00")
        ).replace(tzinfo=None)
    except ValueError:
        pass

    formats = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d %H:%M:%S.%f",
        "%m/%d/%Y %I:%M:%S %p",
        "%d-%m-%Y %H:%M:%S"
    ]

    for fmt in formats:
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue

    return None


# Format timestamps consistently
def normalize_timestamp(value):
    parsed = parse_datetime(value)

    if parsed:
        return parsed.strftime("%Y-%m-%d %H:%M:%S")

    return str(value or "")


# Convert imported log data into a LogEvent object
def event_from_mapping(row, source_name="Imported"):
    row = {
        str(key).strip().lower().replace(" ", "_"): value
        for key, value in row.items()
        if key is not None
    }

    def pick(*names, default=""):
        for name in names:
            value = row.get(name)

            if value not in (None, ""):
                return str(value).strip()

        return default

    return LogEvent(
        timestamp=normalize_timestamp(
            pick("timestamp", "time_created", "time", "date", "datetime")
        ),
        event_type=pick(
            "event_type", "type", "category", default="System Event"
        ),
        source=pick("source", "provider", "log_name", default=source_name),
        description=pick(
            "description", "message", "details", "event",
            default="Imported log entry"
        ),
        host=pick("host", "computer", "computer_name", "hostname"),
        user=pick("user", "username", "account"),
        source_ip=pick("source_ip", "ip", "ip_address", "client_ip"),
        severity=pick("severity", "level", "event_level", default="Info"),
        event_id=pick("event_id", "eventid", "id")
    )


# Read logs from CSV, JSON, TSV, TXT, or EVTX files
def parse_file(path):
    extension = Path(path).suffix.lower()
    events = []

    if extension in (".csv", ".tsv", ".txt"):
        delimiter = "\t" if extension == ".tsv" else ","

        with open(path, "r", encoding="utf-8-sig", newline="") as file:
            for row in csv.DictReader(file, delimiter=delimiter):
                events.append(event_from_mapping(row, Path(path).name))

    elif extension == ".json":
        with open(path, "r", encoding="utf-8-sig") as file:
            data = json.load(file)

        if isinstance(data, dict):
            data = data.get("events", data.get("logs", [data]))

        if not isinstance(data, list):
            raise ValueError("JSON must contain a list of event objects.")

        for row in data:
            if isinstance(row, dict):
                events.append(event_from_mapping(row, Path(path).name))

    elif extension == ".evtx":
        try:
            from Evtx.Evtx import Evtx
        except ImportError:
            raise RuntimeError(
                "Install EVTX support:\n"
                "py -m pip install python-evtx"
            )

        namespace = {
            "e": "http://schemas.microsoft.com/win/2004/08/events/event"
        }

        levels = {
            "1": "Critical",
            "2": "Error",
            "3": "Warning",
            "4": "Info",
            "5": "Verbose"
        }

        with Evtx(path) as log:
            for record in log.records():
                try:
                    root = ET.fromstring(record.xml())
                    system = root.find("e:System", namespace)

                    if system is None:
                        continue

                    provider = system.find("e:Provider", namespace)
                    event_id = system.find("e:EventID", namespace)
                    level = system.find("e:Level", namespace)
                    time_node = system.find("e:TimeCreated", namespace)
                    computer = system.find("e:Computer", namespace)

                    details = []

                    for item in root.findall(
                        ".//e:EventData/e:Data", namespace
                    ):
                        name = item.attrib.get("Name", "")
                        value = item.text or ""
                        details.append(
                            f"{name}={value}" if name else value
                        )

                    events.append(LogEvent(
                        timestamp=normalize_timestamp(
                            time_node.attrib.get("SystemTime", "")
                            if time_node is not None else ""
                        ),
                        event_type="Windows Event",
                        source=(
                            provider.attrib.get("Name", "")
                            if provider is not None else "Windows Event Log"
                        ),
                        description="; ".join(details)
                        or "Windows event record",
                        host=(
                            computer.text
                            if computer is not None and computer.text else ""
                        ),
                        severity=levels.get(
                            level.text if level is not None else "", "Info"
                        ),
                        event_id=(
                            event_id.text
                            if event_id is not None and event_id.text else ""
                        )
                    ))

                except Exception:
                    continue

    else:
        raise ValueError("Unsupported file format.")

    return events


# Control the main application window and its functions
class TimelineApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("1420x850")
        self.root.minsize(1050, 650)
        self.root.configure(bg=BG)

        # Create the linked list and event-processing queue
        self.events = EventLinkedList()
        self.next_id = 1
        self.queue = queue.Queue()

        # Store search, filter, sorting, and status values
        self.search_var = tk.StringVar()
        self.filter_var = tk.StringVar(value="All severities")
        self.sort_key = tk.StringVar(value="timestamp")
        self.sort_order = tk.StringVar(value="Oldest first")
        self.status_var = tk.StringVar(value="Ready")

        self.configure_style()
        self.build_ui()

        # Check for background task results
        self.root.after(120, self.process_queue)

    # Set the appearance of tables and dropdown menus
    def configure_style(self):
        style = ttk.Style()
        style.theme_use("clam")

        style.configure(
            "Treeview",
            background=PANEL,
            foreground=TEXT,
            fieldbackground=PANEL,
            rowheight=30,
            borderwidth=0,
            font=("Segoe UI", 9)
        )

        style.configure(
            "Treeview.Heading",
            background=PANEL2,
            foreground=ACCENT,
            font=("Segoe UI", 9, "bold")
        )

        style.map(
            "Treeview",
            background=[("selected", "#21486b")],
            foreground=[("selected", "#ffffff")]
        )

        style.configure(
            "TCombobox",
            fieldbackground=PANEL2,
            background=PANEL2,
            foreground=TEXT,
            arrowcolor=ACCENT
        )

    # Create a reusable button
    def button(self, parent, text, command, color=PANEL2, fg=TEXT):
        return tk.Button(
            parent,
            text=text,
            command=command,
            bg=color,
            fg=fg,
            activebackground=ACCENT,
            activeforeground=BG,
            relief="flat",
            bd=0,
            padx=12,
            pady=9,
            cursor="hand2",
            font=("Segoe UI", 9, "bold")
        )

    # Display a statistics card
    def stat_card(self, parent, title, color):
        card = tk.Frame(
            parent, bg=PANEL,
            highlightbackground=BORDER, highlightthickness=1
        )
        card.pack(side="left", fill="x", expand=True, padx=(0, 10))

        tk.Label(
            card, text=title, bg=PANEL, fg=MUTED
        ).pack(anchor="w", padx=14, pady=(10, 2))

        value = tk.Label(
            card, text="0", bg=PANEL, fg=color,
            font=("Segoe UI", 18, "bold")
        )
        value.pack(anchor="w", padx=14, pady=(0, 10))

        return value

    # Build the main application interface
    def build_ui(self):
        header = tk.Frame(self.root, bg=BG)
        header.pack(fill="x", padx=24, pady=18)

        tk.Label(
            header, text="FORENSIGHT", bg=BG, fg=ACCENT,
            font=("Segoe UI", 22, "bold")
        ).pack(side="left")

        self.button(
            header, "+ Add Event", self.add_dialog, ACCENT, BG
        ).pack(side="right", padx=5)

        self.button(
            header, "Import Logs", self.import_logs
        ).pack(side="right", padx=5)

        self.button(
            header, "Collect Windows Logs", self.collect_windows_logs
        ).pack(side="right", padx=5)

        # Display event statistics
        stats = tk.Frame(self.root, bg=BG)
        stats.pack(fill="x", padx=24, pady=(0, 14))

        self.total_label = self.stat_card(stats, "TOTAL EVENTS", ACCENT)
        self.alert_label = self.stat_card(stats, "HIGH / ERROR", RED)
        self.source_label = self.stat_card(stats, "SOURCES", PURPLE)
        self.range_label = self.stat_card(stats, "TIME SPAN", GREEN)

        main = tk.Frame(self.root, bg=BG)
        main.pack(fill="both", expand=True, padx=24, pady=(0, 12))

        left = tk.Frame(
            main, bg=PANEL,
            highlightbackground=BORDER, highlightthickness=1
        )
        left.pack(side="left", fill="both", expand=True)

        right = tk.Frame(
            main, bg=PANEL, width=300,
            highlightbackground=BORDER, highlightthickness=1
        )
        right.pack(side="right", fill="y", padx=(14, 0))
        right.pack_propagate(False)

        # Create the search and severity filter controls
        toolbar = tk.Frame(left, bg=PANEL)
        toolbar.pack(fill="x", padx=14, pady=14)

        search = tk.Entry(
            toolbar, textvariable=self.search_var,
            bg=PANEL2, fg=TEXT, insertbackground=ACCENT,
            relief="flat"
        )
        search.pack(side="left", fill="x", expand=True, ipady=9, padx=(0, 10))
        search.bind("<KeyRelease>", lambda event: self.refresh_table())

        filter_box = ttk.Combobox(
            toolbar, textvariable=self.filter_var, state="readonly",
            values=[
                "All severities", "Critical / Error",
                "Warning", "Info / Other"
            ],
            width=16
        )
        filter_box.pack(side="right")
        filter_box.bind(
            "<<ComboboxSelected>>", lambda event: self.refresh_table()
        )

        # Create sorting controls
        sortbar = tk.Frame(left, bg=PANEL)
        sortbar.pack(fill="x", padx=14, pady=(0, 10))

        tk.Label(
            sortbar, text="SORT BY", bg=PANEL, fg=ACCENT
        ).pack(side="left", padx=5)

        sortbox = ttk.Combobox(
            sortbar, textvariable=self.sort_key, state="readonly",
            values=[
                "timestamp", "severity", "event_type",
                "source", "host", "user"
            ],
            width=14
        )
        sortbox.pack(side="left")
        sortbox.bind(
            "<<ComboboxSelected>>", lambda event: self.refresh_table()
        )

        orderbox = ttk.Combobox(
            sortbar, textvariable=self.sort_order, state="readonly",
            values=["Oldest first", "Newest first"], width=14
        )
        orderbox.pack(side="left", padx=6)
        orderbox.bind(
            "<<ComboboxSelected>>", lambda event: self.refresh_table()
        )

        self.button(
            sortbar, "Reset", self.reset_filters
        ).pack(side="left", padx=5)

        self.button(
            sortbar, "Delete Selected", self.delete_selected,
            "#3b2033", "#ffb4c0"
        ).pack(side="right")

        # Display log events in a table
        columns = (
            "timestamp", "severity", "event_type",
            "source", "event_id", "host", "description"
        )

        self.tree = ttk.Treeview(
            left, columns=columns, show="headings", selectmode="extended"
        )

        headings = {
            "timestamp": "TIMESTAMP",
            "severity": "SEVERITY",
            "event_type": "EVENT TYPE",
            "source": "SOURCE",
            "event_id": "EVENT ID",
            "host": "HOST",
            "description": "DESCRIPTION"
        }

        widths = {
            "timestamp": 155,
            "severity": 90,
            "event_type": 125,
            "source": 150,
            "event_id": 75,
            "host": 130,
            "description": 350
        }

        for column in columns:
            self.tree.heading(column, text=headings[column])
            self.tree.column(
                column, width=widths[column],
                minwidth=70, anchor="w"
            )

        self.tree.pack(fill="both", expand=True, padx=14, pady=8)
        self.tree.bind("<<TreeviewSelect>>", self.show_details)

        # Display details of the selected event
        tk.Label(
            right, text="EVENT SUMMARY", bg=PANEL, fg=ACCENT,
            font=("Segoe UI", 11, "bold")
        ).pack(anchor="w", padx=16, pady=15)

        self.details = tk.Text(
            right, bg=PANEL2, fg=TEXT, relief="flat",
            wrap="word", font=("Consolas", 9)
        )
        self.details.pack(fill="both", expand=True, padx=14, pady=8)

        self.set_details("Select an event to view its summary.")

        self.button(
            right, "Create Report", self.create_report, ACCENT, BG
        ).pack(fill="x", padx=14, pady=8)

        self.button(
            right, "Delete All Events", self.clear_all,
            "#3b2033", "#ffb4c0"
        ).pack(fill="x", padx=14, pady=10)

        # Display the application status
        footer = tk.Frame(self.root, bg="#070b16")
        footer.pack(fill="x", side="bottom")

        tk.Label(
            footer, text="● LINKED LIST ACTIVE",
            bg="#070b16", fg=GREEN
        ).pack(side="left", padx=24, pady=9)

        tk.Label(
            footer, textvariable=self.status_var,
            bg="#070b16", fg=MUTED
        ).pack(side="right", padx=24, pady=9)

    # Update the event summary panel
    def set_details(self, text):
        self.details.configure(state="normal")
        self.details.delete("1.0", "end")
        self.details.insert("1.0", text)
        self.details.configure(state="disabled")

    # Insert a new event into the linked list
    def insert_event(self, event, refresh=True):
        event.event_id_internal = self.next_id
        self.next_id += 1
        self.events.append(event)

        if refresh:
            self.refresh_table()

    # Display events based on search, filters, and sorting
    def refresh_table(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        matches = self.events.find(self.search_var.get())
        selected_filter = self.filter_var.get()

        # Group events according to severity
        def severity_group(event):
            severity = event.severity.lower()

            if any(word in severity for word in (
                "critical", "error", "high", "fatal"
            )):
                return "high"

            if any(word in severity for word in (
                "warning", "medium", "suspicious"
            )):
                return "warning"

            return "other"

        # Apply the selected severity filter
        if selected_filter == "Critical / Error":
            matches = [e for e in matches if severity_group(e) == "high"]
        elif selected_filter == "Warning":
            matches = [e for e in matches if severity_group(e) == "warning"]
        elif selected_filter == "Info / Other":
            matches = [e for e in matches if severity_group(e) == "other"]

        # Sort events by the selected field
        key = self.sort_key.get()
        reverse = self.sort_order.get() == "Newest first"

        if key == "timestamp":
            matches.sort(
                key=lambda e: parse_datetime(e.timestamp) or datetime.min,
                reverse=reverse
            )
        else:
            matches.sort(
                key=lambda e: str(getattr(e, key, "")).casefold(),
                reverse=reverse
            )

        # Add the filtered events to the table
        for event in matches:
            self.tree.insert(
                "", "end", iid=str(event.event_id_internal),
                values=(
                    event.timestamp, event.severity, event.event_type,
                    event.source, event.event_id, event.host,
                    event.description
                )
            )

        self.update_stats()

    # Reset search, filter, and sorting options
    def reset_filters(self):
        self.search_var.set("")
        self.filter_var.set("All severities")
        self.sort_key.set("timestamp")
        self.sort_order.set("Oldest first")
        self.refresh_table()

    # Update the statistics shown on the dashboard
    def update_stats(self):
        events = list(self.events.traverse())

        self.total_label.config(text=str(len(events)))

        self.alert_label.config(text=str(sum(
            1 for e in events
            if any(word in e.severity.lower() for word in (
                "critical", "error", "high", "fatal"
            ))
        )))

        self.source_label.config(text=str(len(set(
            e.source for e in events if e.source
        ))))

        times = [parse_datetime(e.timestamp) for e in events]
        times = [t for t in times if t]

        self.range_label.config(
            text=f"{min(times):%b %d} - {max(times):%b %d}"
            if times else "—"
        )

    # Find the selected event in the linked list
    def selected_event(self):
        selection = self.tree.selection()

        if not selection:
            return None

        internal_id = int(selection[0])

        for event in self.events.traverse():
            if event.event_id_internal == internal_id:
                return event

        return None

    # Show the details of the selected event
    def show_details(self, event=None):
        selected = self.selected_event()

        if selected is None:
            return

        self.set_details(
            f"TIMESTAMP\n{selected.timestamp}\n\n"
            f"SEVERITY\n{selected.severity}\n\n"
            f"EVENT TYPE\n{selected.event_type}\n\n"
            f"SOURCE\n{selected.source}\n\n"
            f"EVENT ID\n{selected.event_id}\n\n"
            f"HOST\n{selected.host}\n\n"
            f"USER\n{selected.user}\n\n"
            f"SOURCE IP\n{selected.source_ip}\n\n"
            f"DESCRIPTION\n{selected.description}\n\n"
            f"LINKED LIST NODE\n{selected.event_id_internal}"
        )

    # Open a form to manually add a log event
    def add_dialog(self):
        win = tk.Toplevel(self.root)
        win.title("Add Forensic Event")
        win.geometry("520x760")
        win.minsize(420, 500)
        win.configure(bg=PANEL)
        win.transient(self.root)
        win.grab_set()

        tk.Label(
            win, text="ADD FORENSIC EVENT", bg=PANEL, fg=ACCENT,
            font=("Segoe UI", 15, "bold")
        ).pack(anchor="w", padx=22, pady=16)

        container = tk.Frame(win, bg=PANEL)
        container.pack(fill="both", expand=True, padx=18)

        canvas = tk.Canvas(container, bg=PANEL, highlightthickness=0)
        scrollbar = ttk.Scrollbar(
            container, orient="vertical", command=canvas.yview
        )

        form = tk.Frame(canvas, bg=PANEL)
        form_window = canvas.create_window(
            (0, 0), window=form, anchor="nw"
        )

        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # Adjust the form when the window size changes
        def update_scroll_region(event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))
            canvas.itemconfigure(form_window, width=canvas.winfo_width())

        form.bind("<Configure>", update_scroll_region)
        canvas.bind("<Configure>", update_scroll_region)

        fields = {}

        definitions = [
            ("Timestamp", "timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Event Type", "event_type", "System Event"),
            ("Source", "source", "Manual Entry"),
            ("Severity", "severity", "Info"),
            ("Event ID", "event_id", ""),
            ("Host", "host", ""),
            ("User", "user", ""),
            ("Source IP", "source_ip", "")
        ]

        # Create input fields for event information
        for label, key, default in definitions:
            tk.Label(
                form, text=label, bg=PANEL, fg=MUTED,
                font=("Segoe UI", 9)
            ).pack(anchor="w", pady=(7, 3))

            entry = tk.Entry(
                form, bg=PANEL2, fg=TEXT,
                insertbackground=ACCENT, relief="flat",
                font=("Segoe UI", 10)
            )
            entry.pack(fill="x", ipady=7)
            entry.insert(0, default)
            fields[key] = entry

        tk.Label(
            form, text="Description", bg=PANEL, fg=MUTED,
            font=("Segoe UI", 9)
        ).pack(anchor="w", pady=(10, 3))

        description = tk.Text(
            form, height=5, bg=PANEL2, fg=TEXT,
            insertbackground=ACCENT, relief="flat",
            wrap="word", font=("Segoe UI", 10)
        )
        description.pack(fill="x", pady=(0, 15))

        button_frame = tk.Frame(
            win, bg=PANEL,
            highlightbackground=BORDER, highlightthickness=1
        )
        button_frame.pack(fill="x", side="bottom")

        # Validate and save the new event
        def save_event():
            timestamp = fields["timestamp"].get().strip()

            if not parse_datetime(timestamp):
                messagebox.showerror(
                    "Invalid Timestamp",
                    "Enter a valid timestamp.\n"
                    "Example: 2026-10-11 21:22:14",
                    parent=win
                )
                return

            event_type = fields["event_type"].get().strip()
            event_description = description.get("1.0", "end").strip()

            if not event_type:
                messagebox.showerror(
                    "Missing Information",
                    "Please enter an Event Type.",
                    parent=win
                )
                return

            if not event_description:
                messagebox.showerror(
                    "Missing Information",
                    "Please enter an event Description.",
                    parent=win
                )
                return

            event = LogEvent(
                timestamp=normalize_timestamp(timestamp),
                event_type=event_type,
                source=fields["source"].get().strip(),
                description=event_description,
                severity=fields["severity"].get().strip() or "Info",
                event_id=fields["event_id"].get().strip(),
                host=fields["host"].get().strip(),
                user=fields["user"].get().strip(),
                source_ip=fields["source_ip"].get().strip()
            )

            self.insert_event(event)
            self.status_var.set("Event added successfully")
            win.destroy()

        self.button(
            button_frame, "ADD EVENT", save_event, ACCENT, BG
        ).pack(fill="x", padx=20, pady=12)

        # Enable mouse-wheel scrolling in the form
        def mouse_wheel(event):
            canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        canvas.bind(
            "<Enter>",
            lambda event: canvas.bind_all("<MouseWheel>", mouse_wheel)
        )
        canvas.bind(
            "<Leave>",
            lambda event: canvas.unbind_all("<MouseWheel>")
        )

        fields["timestamp"].focus_set()

    # Select log files to import
    def import_logs(self):
        paths = filedialog.askopenfilenames(
            title="Import Forensic Logs",
            filetypes=[
                ("Supported Logs", "*.csv *.json *.evtx *.tsv *.txt"),
                ("All Files", "*.*")
            ]
        )

        if not paths:
            return

        self.status_var.set("Importing log files...")

        # Read files in a background thread
        def worker():
            imported, errors = [], []

            for path in paths:
                try:
                    imported.extend(parse_file(path))
                except Exception as exc:
                    errors.append(f"{Path(path).name}: {exc}")

            self.queue.put(("import", imported, errors))

        threading.Thread(target=worker, daemon=True).start()

    # Collect recent logs from the local Windows system
    def collect_windows_logs(self):
        if sys.platform != "win32":
            messagebox.showinfo(
                "Windows Only", "This feature requires Windows."
            )
            return

        if not messagebox.askyesno(
            "Collect Windows Logs",
            "Read recent local System, Security and Application logs?\n\n"
            "Security log access may require administrator permission."
        ):
            return

        self.status_var.set("Collecting Windows logs...")

        # Read Windows event logs in a background thread
        def worker():
            try:
                import win32evtlog

                collected, errors = [], []

                for log_type in ("System", "Security", "Application"):
                    handle = None

                    try:
                        handle = win32evtlog.OpenEventLog(
                            "localhost", log_type
                        )

                        flags = (
                            win32evtlog.EVENTLOG_BACKWARDS_READ
                            | win32evtlog.EVENTLOG_SEQUENTIAL_READ
                        )

                        records = win32evtlog.ReadEventLog(
                            handle, flags, 0
                        ) or []

                        for rec in records[:300]:
                            strings = " | ".join(
                                str(value)
                                for value in (rec.StringInserts or [])
                            )

                            severity = {
                                1: "Error",
                                2: "Warning",
                                4: "Info"
                            }.get(rec.EventType, "Info")

                            collected.append(LogEvent(
                                timestamp=normalize_timestamp(
                                    rec.TimeGenerated.Format()
                                ),
                                event_type=log_type,
                                source=str(rec.SourceName or log_type),
                                description=strings or
                                f"Windows event {rec.EventID}",
                                host=str(rec.ComputerName or ""),
                                severity=severity,
                                event_id=str(rec.EventID & 0xFFFF)
                            ))

                    except Exception as exc:
                        errors.append(f"{log_type}: {exc}")

                    finally:
                        if handle is not None:
                            try:
                                win32evtlog.CloseEventLog(handle)
                            except Exception:
                                pass

                self.queue.put(("import", collected, errors))

            except ImportError:
                self.queue.put((
                    "notice",
                    "Install Windows log support:\n"
                    "py -m pip install pywin32"
                ))

        threading.Thread(target=worker, daemon=True).start()

    # Process results returned by background tasks
    def process_queue(self):
        try:
            while True:
                item = self.queue.get_nowait()

                if item[0] == "import":
                    _, imported, errors = item

                    for event in imported:
                        self.insert_event(event, refresh=False)

                    self.refresh_table()
                    self.status_var.set(f"Loaded {len(imported)} events")

                    message = f"Loaded {len(imported)} events."

                    if errors:
                        message += "\n\n" + "\n".join(errors[:8])

                    messagebox.showinfo("Import Complete", message)

                elif item[0] == "notice":
                    messagebox.showwarning("Notice", item[1])

        except queue.Empty:
            pass

        # Check the queue again after a short delay
        self.root.after(120, self.process_queue)

    # Delete selected events from the linked list
    def delete_selected(self):
        selection = self.tree.selection()

        if not selection:
            messagebox.showinfo(
                "Select Event", "Select one or more events first."
            )
            return

        if not messagebox.askyesno(
            "Delete Selected",
            f"Delete {len(selection)} selected events?"
        ):
            return

        for item in selection:
            self.events.delete_by_id(int(item))

        self.refresh_table()
        self.set_details("Select an event to view its summary.")
        self.status_var.set("Selected events deleted")

    # Delete every event from the timeline
    def clear_all(self):
        if not self.events.size:
            return

        if messagebox.askyesno(
            "Delete All Events",
            f"Delete all {self.events.size} events?"
        ):
            self.events.clear()
            self.refresh_table()
            self.set_details("Timeline cleared.")
            self.status_var.set("Timeline cleared")

    # Prepare chronological data for the forensic report
    def get_report_data(self):
        events = sorted(
            list(self.events.traverse()),
            key=lambda event: parse_datetime(event.timestamp) or datetime.min
        )

        # Count high-severity and warning events
        def group(event):
            severity = event.severity.lower()

            if any(word in severity for word in (
                "critical", "error", "high", "fatal"
            )):
                return "high"

            if any(word in severity for word in (
                "warning", "medium", "suspicious"
            )):
                return "warning"

            return "other"

        high = sum(1 for event in events if group(event) == "high")
        warnings = sum(1 for event in events if group(event) == "warning")

        times = [
            parse_datetime(event.timestamp) for event in events
        ]
        times = [value for value in times if value]

        time_range = (
            f"{min(times)} to {max(times)}"
            if times else "Unavailable"
        )

        return events, high, warnings, time_range

    # Open the report format selection window
    def create_report(self):
        if not self.events.size:
            messagebox.showinfo(
                "No Events",
                "Import or collect logs before creating a report."
            )
            return

        win = tk.Toplevel(self.root)
        win.title("Create Forensic Report")
        win.geometry("360x230")
        win.configure(bg=PANEL)
        win.transient(self.root)
        win.grab_set()

        tk.Label(
            win, text="CREATE FORENSIC REPORT",
            bg=PANEL, fg=ACCENT,
            font=("Segoe UI", 13, "bold")
        ).pack(pady=(22, 8))

        tk.Label(
            win, text="Choose your report format",
            bg=PANEL, fg=TEXT
        ).pack(pady=5)

        self.button(
            win, "Word Document (.docx)",
            lambda: self.choose_report(win, "docx")
        ).pack(fill="x", padx=25, pady=5)

        self.button(
            win, "PDF Document (.pdf)",
            lambda: self.choose_report(win, "pdf")
        ).pack(fill="x", padx=25, pady=5)

    # Select and export the chosen report format
    def choose_report(self, window, report_type):
        window.destroy()

        if report_type == "docx":
            self.export_report_docx()
        else:
            self.export_report_pdf()

    # Create a Word forensic report
    def export_report_docx(self):
        try:
            from docx import Document
        except ImportError:
            messagebox.showerror(
                "Missing Package",
                "Run: py -m pip install python-docx"
            )
            return

        path = filedialog.asksaveasfilename(
            title="Save Word Report",
            defaultextension=".docx",
            filetypes=[("Word Document", "*.docx")],
            initialfile="forensic_report.docx"
        )

        if not path:
            return

        events, high, warnings, time_range = self.get_report_data()

        doc = Document()
        doc.add_heading("FORENSIGHT", 0)
        doc.add_paragraph("Digital Forensics Log Timeline Report")
        doc.add_paragraph(
            f"Report generated: {datetime.now():%Y-%m-%d %H:%M:%S}"
        )

        # Add event statistics to the report
        doc.add_heading("Event Summary", level=1)
        doc.add_paragraph(f"Total events: {len(events)}")
        doc.add_paragraph(f"High / Error events: {high}")
        doc.add_paragraph(f"Warning events: {warnings}")
        doc.add_paragraph(
            f"Unique sources: {len(set(e.source for e in events if e.source))}"
        )
        doc.add_paragraph(f"Investigation time range: {time_range}")

        # Create a chronological event table
        doc.add_heading("Chronological Event Timeline", level=1)

        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"

        headings = [
            "Timestamp", "Severity", "Event ID", "Source", "Description"
        ]

        for cell, heading in zip(table.rows[0].cells, headings):
            cell.text = heading

        for event in events:
            cells = table.add_row().cells

            values = [
                event.timestamp, event.severity, event.event_id,
                event.source, event.description
            ]

            for cell, value in zip(cells, values):
                cell.text = str(value)

        # Add investigation notes
        doc.add_heading("Report Notes", level=1)
        doc.add_paragraph(
            "This report summarizes events loaded into the application. "
            "Suspicious patterns require further investigation. "
            "This report alone does not prove that an attack occurred."
        )

        try:
            doc.save(path)
            self.status_var.set("Word report created")
            messagebox.showinfo(
                "Report Created", f"Saved to:\n{path}"
            )
        except Exception as exc:
            messagebox.showerror("Report Error", str(exc))

    # Create a PDF forensic report
    def export_report_pdf(self):
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import (
                SimpleDocTemplate, Paragraph, Spacer,
                Table, TableStyle
            )
        except ImportError:
            messagebox.showerror(
                "Missing Package",
                "Run: py -m pip install reportlab"
            )
            return

        path = filedialog.asksaveasfilename(
            title="Save PDF Report",
            defaultextension=".pdf",
            filetypes=[("PDF Document", "*.pdf")],
            initialfile="forensic_report.pdf"
        )

        if not path:
            return

        events, high, warnings, time_range = self.get_report_data()
        styles = getSampleStyleSheet()

        # Set the PDF page layout
        doc = SimpleDocTemplate(
            path,
            pagesize=landscape(A4),
            rightMargin=10 * mm,
            leftMargin=10 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm
        )

        # Add the report title and event summary
        story = [
            Paragraph("FORENSIGHT", styles["Title"]),
            Paragraph(
                "Digital Forensics Log Timeline Report",
                styles["Heading2"]
            ),
            Paragraph(
                f"Report generated: {datetime.now():%Y-%m-%d %H:%M:%S}",
                styles["Normal"]
            ),
            Spacer(1, 6 * mm),
            Paragraph("Event Summary", styles["Heading1"]),
            Paragraph(f"Total events: {len(events)}", styles["Normal"]),
            Paragraph(f"High / Error events: {high}", styles["Normal"]),
            Paragraph(f"Warning events: {warnings}", styles["Normal"]),
            Paragraph(
                f"Unique sources: {len(set(e.source for e in events if e.source))}",
                styles["Normal"]
            ),
            Paragraph(
                f"Investigation time range: {time_range}",
                styles["Normal"]
            ),
            Spacer(1, 6 * mm),
            Paragraph(
                "Chronological Event Timeline",
                styles["Heading1"]
            )
        ]

        # Prepare event details for the PDF table
        data = [[
            "Timestamp", "Severity", "Event ID", "Source", "Description"
        ]]

        for event in events:
            data.append([
                Paragraph(str(event.timestamp), styles["BodyText"]),
                Paragraph(str(event.severity), styles["BodyText"]),
                Paragraph(str(event.event_id), styles["BodyText"]),
                Paragraph(str(event.source), styles["BodyText"]),
                Paragraph(str(event.description), styles["BodyText"])
            ])

        table = Table(
            data,
            colWidths=[
                36 * mm, 23 * mm, 19 * mm, 40 * mm, 145 * mm
            ],
            repeatRows=1
        )

        # Format the PDF table
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17233b")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5)
        ]))

        story.extend([
            table,
            Spacer(1, 6 * mm),
            Paragraph("Report Notes", styles["Heading1"]),
            Paragraph(
                "This report summarizes available events. Further "
                "investigation is required before drawing conclusions "
                "about a possible incident.",
                styles["Normal"]
            )
        ])

        # Save the completed PDF report
        try:
            doc.build(story)
            self.status_var.set("PDF report created")
            messagebox.showinfo(
                "Report Created", f"Saved to:\n{path}"
            )
        except Exception as exc:
            messagebox.showerror("Report Error", str(exc))


# Start the Digital Forensics Log Timeline Builder
def main():
    root = tk.Tk()
    TimelineApp(root)
    root.mainloop()


# Run the application
if __name__ == "__main__":
    main()